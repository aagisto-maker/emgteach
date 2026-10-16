"""Convierte neurotransmision_marp.md en un PPTX nativo y editable.

Reproduce el estilo del Markdown de Marp (fondo #0f172a, tarjetas, etiquetas
de color, línea temporal y tabla) con formas y cuadros de texto de PowerPoint,
y pasa los comentarios HTML de cada diapositiva a las notas del orador.

Uso:  python marp_a_pptx_editable.py [entrada.md] [salida.pptx]
"""
import math
import re
import sys

from bs4 import BeautifulSoup, NavigableString, Tag
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

# Lienzo de Marp: 1280 x 720 px. Todas las medidas del script van en px.
PX = 9525  # EMU por px


def px(v):
    return Emu(int(round(v * PX)))


def fpt(v):
    """Tamaño de letra: px de CSS -> pt."""
    return Pt(v * 0.75)


def rgb(hexstr):
    return RGBColor.from_string(hexstr.lstrip("#").upper())


BG = "0f172a"
TEXT = "e2e8f0"
H1 = "38bdf8"
H2 = "94a3b8"
CARD = "1e293b"
BORDER = "334155"
HIGHLIGHT = "38bdf8"
BADGES = {"badge-blue": "0284c7", "badge-amber": "d97706",
          "badge-green": "059669", "badge-red": "dc2626"}
FONT = "Segoe UI"

W, H = 1280, 720
PAD_X, PAD_TOP, PAD_BOTTOM = 65, 45, 50
H1_PX, H2_PX = 40, 26
LINE = 1.3  # interlineado relativo


# ---------------------------------------------------------------- lectura .md

def parse_inline_md(text):
    """'**negrita**' y '*cursiva*' -> lista de (texto, negrita, cursiva)."""
    runs = []
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not tok:
            continue
        if tok.startswith("**"):
            runs.append((tok[2:-2], True, False))
        elif tok.startswith("*"):
            runs.append((tok[1:-1], False, True))
        else:
            runs.append((tok, False, False))
    return runs


def html_runs(node, bold=False, italic=False):
    runs = []
    for child in node.children:
        if isinstance(child, NavigableString):
            txt = re.sub(r"\s+", " ", str(child))
            if txt:
                runs.append((txt, bold, italic))
        elif isinstance(child, Tag):
            if child.name == "br":
                runs.append(("\n", bold, italic))
            else:
                runs.extend(html_runs(child, bold or child.name in ("b", "strong"),
                                      italic or child.name in ("i", "em")))
    return clean_runs(runs)


def clean_runs(runs):
    """Quita espacios sobrantes al principio y al final del párrafo."""
    while runs and not runs[0][0].strip():
        runs.pop(0)
    while runs and not runs[-1][0].strip():
        runs.pop()
    if runs:
        runs[0] = (runs[0][0].lstrip(), *runs[0][1:])
        runs[-1] = (runs[-1][0].rstrip(), *runs[-1][1:])
    return runs


def parse_card(div):
    card = {"highlight": "card-highlight" in div.get("class", []),
            "badge": None, "paras": []}
    for child in div.children:
        if not isinstance(child, Tag):
            continue
        classes = child.get("class", [])
        if "badge" in classes:
            color = next((BADGES[c] for c in classes if c in BADGES), BADGES["badge-blue"])
            card["badge"] = (child.get_text(strip=True), color)
        else:
            runs = html_runs(child)
            if runs:
                card["paras"].append(runs)
    return card


def parse_slide(raw):
    slide = {"lead": False, "notes": "", "h1": "", "h2": "", "blocks": []}
    for comment in re.findall(r"<!--(.*?)-->", raw, flags=re.S):
        c = comment.strip()
        if re.match(r"^_?\w+\s*:", c) and "\n" not in c:
            if "lead" in c:
                slide["lead"] = True
        else:
            slide["notes"] = re.sub(r"^NOTAS DEL ORADOR:\s*", "", c).strip()
    body = re.sub(r"<!--.*?-->", "", raw, flags=re.S)

    html_buf, table, text_lines = [], [], []
    depth = 0
    for line in body.splitlines():
        s = line.strip()
        if depth or s.startswith("<div"):
            html_buf.append(line)
            depth += line.count("<div") - line.count("</div")
            if depth == 0:
                soup = BeautifulSoup("\n".join(html_buf), "html.parser")
                top = soup.find("div")
                classes = top.get("class", [])
                if "timeline" in classes:
                    steps = [html_runs(d) for d in top.find_all("div", class_="timeline-step")]
                    slide["blocks"].append({"type": "timeline", "steps": steps})
                elif "grid-2" in classes or "grid-3" in classes:
                    cards = [parse_card(d) for d in top.find_all("div", recursive=False)]
                    slide["blocks"].append({"type": "grid", "cards": cards})
                else:
                    slide["blocks"].append({"type": "grid", "cards": [parse_card(top)],
                                            "gap_top": 20})
                html_buf = []
        elif s.startswith("# "):
            slide["h1"] = s[2:]
        elif s.startswith("## "):
            slide["h2"] = s[3:]
        elif s.startswith("|"):
            cells = [c.strip() for c in s.strip("|").split("|")]
            if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                table.append([parse_inline_md(c) for c in cells])
        elif s:
            text_lines.append(parse_inline_md(s))
    if table:
        slide["blocks"].append({"type": "table", "rows": table})
    if text_lines:
        slide["blocks"].insert(0, {"type": "text", "paras": text_lines})
    return slide


def read_marp(path):
    content = open(path, encoding="utf-8").read()
    content = re.sub(r"^---\n.*?\n---\n", "", content, count=1, flags=re.S)
    return [parse_slide(chunk) for chunk in re.split(r"\n---\n", content) if chunk.strip()]


# ---------------------------------------------------------------- dibujo

def set_lang(run):
    run._r.get_or_add_rPr().set("lang", "es-ES")


def add_runs(paragraph, runs, size_px, color=TEXT, bold=False):
    for text, b, i in runs:
        r = paragraph.add_run()
        r.text = text
        r.font.name = FONT
        r.font.size = fpt(size_px)
        r.font.bold = bold or b
        r.font.italic = i
        r.font.color.rgb = rgb(color)
        set_lang(r)


def textbox(slide, x, y, w, h, paras, size_px, color=TEXT, bold=False,
            align=PP_ALIGN.LEFT, space_after=0.45, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for k, runs in enumerate(paras):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = 1.1
        if k < len(paras) - 1:
            p.space_after = fpt(size_px * space_after)
        add_runs(p, runs, size_px, color, bold)
    return tb


def rect(slide, x, y, w, h, fill, line=None, line_px=1, radius_px=0):
    shape = MSO_SHAPE.ROUNDED_RECTANGLE if radius_px else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape, px(x), px(y), px(w), px(h))
    if radius_px:
        s.adjustments[0] = min(0.5, radius_px / min(w, h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_px * 0.75)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def text_height(paras, width, size_px, space_after=0.45):
    """Altura estimada (px) de un bloque de párrafos con ajuste de línea."""
    chars_per_line = max(1, width / (size_px * 0.53))
    lines = 0
    for runs in paras:
        for chunk in "".join(t for t, _, _ in runs).split("\n"):
            lines += max(1, math.ceil(len(chunk) / chars_per_line))
    return lines * size_px * LINE + (len(paras) - 1) * size_px * space_after


def card_height(card, width, size_px):
    h = 36  # padding vertical 18 + 18
    if card["badge"]:
        h += 38
    return h + text_height(card["paras"], width - 44, size_px)


def draw_card(slide, x, y, w, h, card, size_px):
    if card["highlight"]:
        rect(slide, x, y, w, h, CARD, HIGHLIGHT, 2, radius_px=12)
    else:
        rect(slide, x, y, w, h, CARD, BORDER, 1, radius_px=12)
    cy = y + 18
    if card["badge"]:
        label, color = card["badge"]
        label = label.upper()
        bw = len(label) * 10.5 + 24
        b = rect(slide, x + 22, cy, bw, 28, color, radius_px=6)
        tf = b.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        add_runs(p, [(label, True, False)], 16, "ffffff")
        cy += 38
    textbox(slide, x + 22, cy, w - 44, y + h - cy - 12, card["paras"], size_px)


def draw_grid(slide, block, top, bottom):
    cards = block["cards"]
    n = len(cards)
    gap = 24 if n == 2 else 18
    top += block.get("gap_top", 15)
    cw = (W - 2 * PAD_X - gap * (n - 1)) / n
    size = 28
    while size > 16 and max(card_height(c, cw, size) for c in cards) > bottom - top:
        size -= 1
    ch = max(card_height(c, cw, size) for c in cards)
    for i, c in enumerate(cards):
        draw_card(slide, PAD_X + i * (cw + gap), top, cw, ch, c, size)
    return top + ch


def draw_timeline(slide, block, top):
    steps = block["steps"]
    n, gap = len(steps), 12
    top += 25
    sw = (W - 2 * PAD_X - gap * (n - 1)) / n
    size = 21
    # Primera línea de cada paso en negrita (es el <b> del Markdown)
    paras = [[[r for r in s if r[1]], [r for r in s if not r[1] and r[0] != "\n"]] for s in steps]
    paras = [[clean_runs(p) for p in ps] for ps in paras]
    sh = max(text_height(p, sw - 28, size, 0.2) for p in paras) + 44
    for i, p in enumerate(paras):
        x = PAD_X + i * (sw + gap)
        rect(slide, x, top, sw, sh, CARD, radius_px=8)
        rect(slide, x, top, sw, 4, HIGHLIGHT)
        textbox(slide, x + 14, top + 16, sw - 28, sh - 24, p, size,
                align=PP_ALIGN.CENTER, space_after=0.2)
    return top + sh


def set_cell_border(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = etree.SubElement(tcPr, qn(tag), w=str(int(Pt(0.75))), cap="flat", cmpd="sng", algn="ctr")
        fill = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(fill, qn("a:srgbClr"), val=color)
        etree.SubElement(ln, qn("a:prstDash"), val="solid")


def draw_table(slide, block, top):
    rows = block["rows"]
    n_rows, n_cols = len(rows), len(rows[0])
    top += 10
    widths = [0.18, 0.24, 0.26, 0.32]
    if len(widths) != n_cols:
        widths = [1 / n_cols] * n_cols
    tw = W - 2 * PAD_X
    size = 22
    row_h = [max(text_height([c], tw * widths[j] - 20, size) for j, c in enumerate(r)) + 20
             for r in rows]
    shape = slide.shapes.add_table(n_rows, n_cols, px(PAD_X), px(top), px(tw), px(sum(row_h)))
    tbl = shape.table
    tbl.first_row = tbl.horz_banding = False
    tblPr = tbl._tbl.tblPr
    style = tblPr.find(qn("a:tableStyleId"))
    if style is not None:
        tblPr.remove(style)
    for j, f in enumerate(widths):
        tbl.columns[j].width = px(tw * f)
    for i, r in enumerate(rows):
        tbl.rows[i].height = px(row_h[i])
        for j, runs in enumerate(r):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(CARD if i == 0 else BG)
            cell.margin_left = cell.margin_right = cell.margin_top = cell.margin_bottom = px(10)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            if i == 0:
                add_runs(p, runs, size, H1, bold=True)
            else:
                add_runs(p, runs, size)
            set_cell_border(cell, BORDER)
    return top + sum(row_h)


def draw_heading(slide, data, top):
    textbox(slide, PAD_X, top, W - 2 * PAD_X, 56, [[(data["h1"], True, False)]], H1_PX, H1)
    top += 60
    rect(slide, PAD_X, top, W - 2 * PAD_X, 2, BORDER)
    top += 16
    if data["h2"]:
        textbox(slide, PAD_X, top, W - 2 * PAD_X, 36, [[(data["h2"], False, False)]], H2_PX, H2)
        top += 42
    return top


def build_slide(prs, data, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(BG)

    if data["lead"]:
        paras = data["blocks"][0]["paras"] if data["blocks"] else []
        body_h = text_height(paras, W - 2 * PAD_X, 24, 0.1)
        top = (H - (60 + 18 + 42 + 20 + body_h)) / 2
        top = draw_heading(slide, data, top) + 20
        textbox(slide, PAD_X, top, W - 2 * PAD_X, body_h + 10, paras, 24, space_after=0.1)
    else:
        top = draw_heading(slide, data, PAD_TOP)
        for block in data["blocks"]:
            if block["type"] == "grid":
                top = draw_grid(slide, block, top, H - PAD_BOTTOM)
            elif block["type"] == "timeline":
                top = draw_timeline(slide, block, top)
            elif block["type"] == "table":
                top = draw_table(slide, block, top)

    textbox(slide, W - PAD_X - 60, H - 34, 60, 20, [[(str(number), False, False)]],
            16, H2, align=PP_ALIGN.RIGHT)
    if data["notes"]:
        slide.notes_slide.notes_text_frame.text = data["notes"]


def build(md_path, out_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = px(W), px(H)
    for i, data in enumerate(read_marp(md_path), start=1):
        build_slide(prs, data, i)
    prs.save(out_path)
    return out_path


if __name__ == "__main__":
    md = sys.argv[1] if len(sys.argv) > 1 else "neurotransmision_marp.md"
    out = sys.argv[2] if len(sys.argv) > 2 else "Neurotransmision_Sinaptica_Marp_Editable.pptx"
    print(build(md, out))
