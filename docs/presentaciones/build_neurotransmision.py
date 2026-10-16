"""Genera la presentación «Neurotransmisión sináptica» (16:9, tema oscuro).

Uso:  python build_neurotransmision.py [salida.pptx]
"""
import math
import sys

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

# Paleta cromática
BG_COLOR = RGBColor(11, 19, 43)        # #0B132B (Azul oscuro uniforme)
CARD_BG = RGBColor(19, 32, 67)         # #132043 (Contenedores)
TEXT_WHITE = RGBColor(240, 244, 248)   # Blanco primario
TEXT_MUTED = RGBColor(170, 190, 215)   # Gris cian secundario
ACCENT_CYAN = RGBColor(0, 180, 216)    # #00B4D8
ACCENT_ORANGE = RGBColor(247, 127, 0)  # #F77F00
ACCENT_GREEN = RGBColor(6, 214, 160)   # #06D6A0
ACCENT_PURPLE = RGBColor(163, 116, 255)  # #A374FF

FONT = "Calibri"

slides_data = [
    # 1. Caso Clínico
    {
        "category": "CASO CLÍNICO INICIAL",
        "title": "La Paciente que se Cansa de Mirar",
        "cards": [
            {
                "title": "Presentación de la Paciente",
                "color": ACCENT_CYAN,
                "items": [
                    ("Perfil:", "Mujer de 28 años sin antecedentes médicos relevantes."),
                    ("Motivo de consulta:", "Debilidad muscular fluctuante y fatigable."),
                    ("Patrón temporal:", "Empeora notablemente hacia la tarde y tras el uso repetido; mejora tras el reposo."),
                ],
            },
            {
                "title": "Signos Cardinales",
                "color": ACCENT_ORANGE,
                "items": [
                    ("Ptosis palpebral:", "Caída bilateral asimétrica de párpados tras fijar la mirada hacia arriba."),
                    ("Diplopía vespertina:", "Visión doble fluctuante que interrumpe actividades cotidianas."),
                    ("Sin déficit sensitivo:", "Reflejos osteotendinosos y sensibilidad conservados."),
                ],
            },
            {
                "title": "La Pregunta Central de la Fisiopatología",
                "color": ACCENT_GREEN,
                "items": [
                    ("¿El cable (nervio)?:", "¿Falla la conducción axonal o la generación del potencial de acción?"),
                    ("¿El motor (músculo)?:", "¿Falla la maquinaria contráctil de la fibra muscular?"),
                    ("¿La interfaz (sinapsis)?:", "¿O claudica la maquinaria química que comunica a ambos en la placa motora?"),
                ],
            },
        ],
        "notes": "Empecemos en la consulta. Nos llega una paciente joven cuyo motivo de consulta parece sutil pero es revelador: 'por la tarde se me caen los párpados y veo doble; por la mañana amanezco bien'. Esta fatigabilidad fluctuante nos sitúa en una encrucijada fisiopatológica clásica: ¿está fallando la conducción axonal, la maquinaria contráctil de la fibra muscular o el sistema de acoplamiento químico intermedio? A lo largo de esta sesión desgranaremos la arquitectura de la neurotransmisión sináptica para descubrir con precisión quirúrgica qué componente molecular de esta paciente ha dejado de funcionar.",
    },
    # 2. Eléctrica vs Química
    {
        "category": "INFRAESTRUCTURA DE LA COMUNICACIÓN",
        "title": "Sinapsis Eléctrica vs. Sinapsis Química",
        "cards": [
            {
                "title": "Sinapsis Eléctrica (El Puente Directo)",
                "color": ACCENT_GREEN,
                "items": [
                    ("Estructura:", "Uniones gap formadas por hemicanales de conexinas (conexones) alineados."),
                    ("Cinética:", "Cuasi-instantánea (sin retardo; ~0.2 ms)."),
                    ("Dirección:", "Bidireccional (acoplamiento electrotónico directo)."),
                    ("Función biológica:", "Sincronización masiva de poblaciones celulares (interneuronas, miocardio, músculo liso)."),
                ],
            },
            {
                "title": "Sinapsis Química (La Computadora Molecular)",
                "color": ACCENT_CYAN,
                "items": [
                    ("Estructura:", "Hendidura sináptica (20-40 nm), vesículas presinápticas y densidad postsináptica."),
                    ("Cinética:", "Retardo sináptico obligado (0.5 a 1.0 ms) por exocitosis y difusión."),
                    ("Dirección:", "Estrictamente unidireccional (presinapsis → postsinapsis)."),
                    ("Función biológica:", "Plasticidad, amplificación de señal, modulación graduada y principal diana farmacológica."),
                ],
            },
        ],
        "notes": "Para entender la patología sináptica, primero debemos diferenciar las dos infraestructuras de comunicación del sistema nervioso. La sinapsis eléctrica es un puente puramente mecánico: canales de conexinas que unen citoplasmas directamente; no hay mensajeros, solo sincronización inmediata en bloque. En cambio, la sinapsis química paga un peaje de tiempo —el retardo sináptico de casi un milisegundo— a cambio de versatilidad: amplifica la señal, permite divergencia computacional y, lo más importante para la medicina, es plástica y modificable mediante fármacos.",
    },
    # 3. Ciclo Sináptico
    {
        "category": "MECANISMO GENERAL",
        "title": "El Ciclo Sináptico y los Criterios de un Mensajero",
        "cards": [
            {
                "title": "El Manifiesto del Neurotransmisor",
                "color": ACCENT_ORANGE,
                "items": [
                    ("1. Síntesis:", "Presencia y síntesis confirmada en la neurona presináptica."),
                    ("2. Liberación:", "Liberación presináptica dependiente de despolarización y Ca2+."),
                    ("3. Identidad de acción:", "La aplicación exógena mimetiza exactamente el efecto fisiológico."),
                    ("4. Terminación:", "Existencia de un mecanismo específico de eliminación o degradación."),
                    ("Nota sobre gases:", "Gases como el NO rompen estas reglas al no empaquetarse ni recaptarse."),
                ],
            },
            {
                "title": "Las 5 Fases del Ciclo Sináptico",
                "color": ACCENT_CYAN,
                "items": [
                    ("1. Síntesis:", "A partir de precursores específicos mediante enzimas en axón o soma."),
                    ("2. Empaquetamiento:", "Bombas H+-ATPasa generan el gradiente para transportadores (VMAT, VAChT)."),
                    ("3. Liberación:", "Apertura de canales de Ca2+ voltaje-dependientes gatilla la exocitosis."),
                    ("4. Acción:", "Unión a receptores postsinápticos específicos (ionotrópicos o metabotrópicos)."),
                    ("5. Inactivación:", "Limpieza de la hendidura por recaptación presináptica/glial o degradación."),
                ],
            },
        ],
        "notes": "La transmisión química no es un evento estático, sino una cadena de montaje en 5 fases continuas. Fíjense en el paso 2: llenar una vesícula requiere energía activa generada por bombas de protones para que transportadores como VMAT o VAChT metan el transmisor contra gradiente. En el paso 3, el calcio actúa como el gatillo indispensable. Y en el paso 5, la señal debe apagarse de inmediato; si el transmisor se queda flotando en la hendidura, la neurona postsináptica se satura o se desensibiliza. Intervenir en cualquiera de estas cinco fases es el fundamento de la neurofarmacología moderna.",
    },
    # 4. SNARE y Toxinas
    {
        "category": "MAQUINARIA BIOFÍSICA",
        "title": "La Fusión Sináptica: Complejo SNARE y Neurotoxinas",
        "cards": [
            {
                "title": "La Cremallera Molecular (Complejo SNARE)",
                "color": ACCENT_CYAN,
                "items": [
                    ("v-SNARE:", "Sinaptobrevina (VAMP) anclada firmemente a la vesícula."),
                    ("t-SNARE:", "Sintaxina-1 y SNAP-25 ancladas en la membrana plasmática diana."),
                    ("El Freno y Gatillo:", "Sinaptotagmina-1 actúa de sensor de Ca2+; sujeta el proceso hasta la entrada del ion y gatilla la fusión en < 1 ms."),
                    ("Mecánica:", "El enrollamiento en espiral de las cuatro hélices tracciona las membranas hasta vencer su repulsión electrostática."),
                ],
            },
            {
                "title": "Las Tijeras Clostridiales (Toxinas)",
                "color": ACCENT_ORANGE,
                "items": [
                    ("Toxina Botulínica:", "Escinde SNAP-25, Sintaxina y Sinaptobrevina en la terminal colinérgica periférica → Bloqueo de liberación de ACh → Parálisis flácida."),
                    ("Toxina Tetánica:", "Escinde Sinaptobrevina tras transporte retrógrado hacia la médula espinal."),
                    ("Mecanismo en Tétanos:", "Bloquea exocitosis de GABA y glicina en interneuronas de Renshaw inhibitorias → Desinhibición motora masiva → Parálisis espástica extrema."),
                ],
            },
        ],
        "notes": "Las dos membranas celulares tienen cargas negativas y se repelen con fuerza. Para fusionarlas hace falta una auténtica prensa mecánica: el complejo SNARE. Sinaptobrevina en la vesícula se enrolla en espiral con SNAP-25 y Sintaxina en la membrana presináptica, tensando la cuerda como una cremallera. Sinaptotagmina actúa de freno de seguridad hasta que detecta el calcio entrante y lo libera. Las toxinas de Clostridium son tijeras moleculares específicas contra este complejo: Botulino corta a nivel periférico causando parálisis flácida; Tétanos viaja retrógradamente por el axón hasta las interneuronas inhibitorias medulares, silenciando el freno gabaérgico y desatando espasmos incontrolables.",
    },
    # 5. Ionotrópicos vs Metabotrópicos
    {
        "category": "RECEPTORES Y TRANSDUCCIÓN",
        "title": "Decodificando el Mensaje: Ionotrópicos vs. Metabotrópicos",
        "cards": [
            {
                "title": "Receptor Ionotrópico (El Canal Rápido)",
                "color": ACCENT_CYAN,
                "items": [
                    ("Mecanismo:", "Poro iónico intrínseco directo abierto por la fijación del ligando."),
                    ("Cinética:", "Respuesta en milisegundos (1-10 ms); efecto rápido, breve y fásico."),
                    ("Ejemplos clave:", "Receptor Nicotínico de ACh (Na+/Ca2+), GABAA (Cl-), receptores glutamatérgicos AMPA y NMDA."),
                    ("Función:", "Transmisión excitatoria e inhibitoria rápida punto a punto."),
                ],
            },
            {
                "title": "Receptor Metabotrópico (La Máquina de Cascada)",
                "color": ACCENT_PURPLE,
                "items": [
                    ("Mecanismo:", "Acoplado a proteínas G heterotriméricas (Gs, Gi, Gq) y cascadas enzimáticas."),
                    ("Cinética:", "Respuesta lenta (segundos a minutos); duradera y altamente amplificada."),
                    ("Ejemplos clave:", "Receptores Muscarínicos (M1-M5), Adrenérgicos (alfa, beta), Dopaminérgicos (D1-D5), GABAB."),
                    ("Función:", "Modulación plástica, fosforilación celular y regulación génica."),
                ],
            },
            {
                "title": "El Termostato Local Presináptico",
                "color": ACCENT_GREEN,
                "items": [
                    ("Autorreceptores:", "Detectan el neurotransmisor recién liberado por la propia terminal y frenan la exocitosis subsiguiente (retroalimentación negativa)."),
                    ("Heterorreceptores:", "Permiten que transmisores de sinapsis vecinas modulen la liberación local de forma cruzada."),
                ],
            },
        ],
        "notes": "Una vez en la hendidura, el mensaje adquiere significado según el receptor que lo reciba. Los ionotrópicos son interruptores directos: se une la molécula, el canal se abre y fluyen iones en cuestión de milisegundos; controlan la neurotransmisión rápida punto a punto. Los metabotrópicos son computadoras celulares: reclutan proteínas G y cascadas bioquímicas complejas. No solo modifican el potencial de membrana, sino que fosforilan dianas y alteran la expresión génica. Noten además los receptores presinápticos: el autorreceptor es el termostato de la neurona para evitar la hiperestimulación apagando la exocitosis.",
    },
    # 6. Circuitos Focales vs Difusos
    {
        "category": "ARQUITECTURA DE CIRCUITOS",
        "title": "Anatomía Sináptica: Sistemas Focales vs. Sistemas de Proyección Difusa",
        "cards": [
            {
                "title": "Sistemas Cableados (Punto a Punto)",
                "color": ACCENT_CYAN,
                "items": [
                    ("Mediadores principales:", "Glutamato (excitador primario del SNC) y GABA (inhibidor primario)."),
                    ("Arquitectura de red:", "Circuitos con conexiones topográficas sumamente precisas, discretas y milimétricas."),
                    ("Función de procesamiento:", "Transmisión de información sensorial discriminativa, control motor fino y computación cortical analítica."),
                ],
            },
            {
                "title": "Sistemas de Radiación Difusa (Sistemas Moduladores)",
                "color": ACCENT_PURPLE,
                "items": [
                    ("Mediadores principales:", "Monoaminas (noradrenalina, serotonina, dopamina, histamina) y acetilcolina central."),
                    ("Arquitectura de red:", "Pocos miles de neuronas en núcleos del tronco y prosencéfalo basal (locus coeruleus, rafe, sustancia negra/ATV, núcleo basal de Meynert) con axones enormemente ramificados."),
                    ("Modo de liberación:", "Varicosidades axonales y transmisión de volumen sobre receptores mayoritariamente metabotrópicos."),
                    ("Función de procesamiento:", "Ajustan el «tono» global: vigilia, atención, ánimo, motivación y recompensa; dianas de antidepresivos y antipsicóticos."),
                ],
            },
        ],
        "notes": "Si miramos el cerebro a escala de circuito, conviven dos filosofías de cableado. Los sistemas focales de glutamato y GABA son como una red telefónica: cada conexión lleva un mensaje concreto a un destinatario concreto, con precisión milimétrica y en milisegundos. Los sistemas difusos son, en cambio, como una emisora de radio: unos pocos núcleos del tronco del encéfalo y del prosencéfalo basal envían axones que se ramifican por toda la corteza y liberan monoaminas o acetilcolina desde varicosidades, actuando sobre receptores metabotrópicos. No transmiten el contenido del mensaje, sino que ajustan el volumen y el estado del sistema: vigilia, atención, ánimo o motivación. Por eso la mayoría de los psicofármacos actúan precisamente sobre estos sistemas.",
    },
    # 7. Resolución del caso
    {
        "category": "RESOLUCIÓN DEL CASO CLÍNICO",
        "title": "Miastenia Gravis: Cuando Falla el Receptor Postsináptico",
        "cards": [
            {
                "title": "El Mecanismo Molecular",
                "color": ACCENT_ORANGE,
                "items": [
                    ("Diana autoinmune:", "Anticuerpos contra el receptor nicotínico de ACh (≈85%) o contra MuSK/LRP4 en la placa motora."),
                    ("Lesión:", "Bloqueo, internalización acelerada y destrucción del receptor mediada por complemento; simplificación de los pliegues postsinápticos."),
                    ("Consecuencia:", "El potencial de placa reduce su amplitud y se pierde el factor de seguridad de la transmisión."),
                ],
            },
            {
                "title": "Por Qué se Fatiga",
                "color": ACCENT_CYAN,
                "items": [
                    ("Depleción fisiológica:", "Con la estimulación repetida, la ACh liberada por impulso disminuye de forma normal."),
                    ("Sin margen:", "Con menos receptores, el potencial de placa cae por debajo del umbral y fibras musculares dejan de responder."),
                    ("Músculos oculares:", "Su alta frecuencia de descarga y menor factor de seguridad explican la ptosis y la diplopía."),
                ],
            },
            {
                "title": "Confirmación y Tratamiento",
                "color": ACCENT_GREEN,
                "items": [
                    ("Estimulación repetitiva (3 Hz):", "Respuesta decremental >10% del potencial de acción muscular compuesto."),
                    ("EMG de fibra aislada:", "Jitter aumentado y bloqueos de transmisión."),
                    ("Tratamiento:", "Inhibidores de acetilcolinesterasa (piridostigmina), inmunosupresión y timectomía cuando procede."),
                ],
            },
        ],
        "notes": "Volvamos a nuestra paciente. El nervio conduce bien y el músculo es capaz de contraerse: lo que falla es la interfaz. Sus anticuerpos atacan el receptor nicotínico de la placa motora, de modo que cada molécula de acetilcolina encuentra menos receptores disponibles. En reposo todavía hay margen, pero con el uso repetido la liberación de acetilcolina por impulso disminuye fisiológicamente y el potencial de placa cae por debajo del umbral: aparece la fatiga. Los músculos oculares, con menor factor de seguridad, son los primeros en claudicar. La estimulación nerviosa repetitiva a 3 Hz muestra el decremento característico, y el tratamiento sintomático consiste precisamente en intervenir la fase 5 del ciclo sináptico: inhibir la acetilcolinesterasa para que la acetilcolina permanezca más tiempo en la hendidura.",
    },
    # 8. Diagnóstico diferencial
    {
        "category": "SÍNTESIS INTEGRADORA",
        "title": "Patología de la Placa Motora: Dónde Falla Cada Enfermedad",
        "cards": [
            {
                "title": "Miastenia Gravis (Postsináptica)",
                "color": ACCENT_CYAN,
                "items": [
                    ("Diana:", "Receptor nicotínico de ACh (autoanticuerpos)."),
                    ("Clínica:", "Debilidad fatigable de predominio ocular y bulbar."),
                    ("Estimulación repetitiva:", "Decremento a baja frecuencia; sin incremento significativo tras el ejercicio."),
                ],
            },
            {
                "title": "Síndrome de Lambert-Eaton (Presináptico)",
                "color": ACCENT_PURPLE,
                "items": [
                    ("Diana:", "Canales de Ca2+ voltaje-dependientes tipo P/Q; frecuente asociación paraneoplásica (carcinoma microcítico)."),
                    ("Clínica:", "Debilidad proximal, arreflexia y disautonomía; mejora transitoria con el esfuerzo."),
                    ("Estimulación repetitiva:", "PAMC basal bajo con incremento >100% a alta frecuencia o tras ejercicio breve."),
                ],
            },
            {
                "title": "Botulismo (Presináptico)",
                "color": ACCENT_ORANGE,
                "items": [
                    ("Diana:", "Proteínas SNARE escindidas por la toxina botulínica."),
                    ("Clínica:", "Parálisis flácida descendente con afectación de pares craneales y pupilas midriáticas."),
                    ("Estimulación repetitiva:", "PAMC bajo con incremento moderado a alta frecuencia."),
                ],
            },
        ],
        "notes": "Cerramos integrando todo lo visto. Cada fase del ciclo sináptico tiene su enfermedad. En la miastenia gravis la presinapsis libera acetilcolina con normalidad, pero el receptor está mermado: decremento a baja frecuencia. En el síndrome de Lambert-Eaton el problema es la entrada de calcio en la terminal: se libera poca acetilcolina por impulso, pero la estimulación de alta frecuencia acumula calcio y facilita la liberación, de ahí el incremento. En el botulismo la toxina corta la maquinaria SNARE y la exocitosis se bloquea. El electromiograma con estimulación repetitiva nos permite, por tanto, localizar el fallo en la sinapsis con una precisión que la clínica sola no alcanza.",
    },
]


def set_background(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE):
    shp = slide.shapes.add_shape(shape, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def style_run(run, size, color, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    rpr = run._r.get_or_add_rPr()
    rpr.set("lang", "es-ES")


def add_textbox(slide, x, y, w, h, text, size, color, bold=False, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    run = p.add_run()
    style_run(run, size, color, bold)
    run.text = text
    return tb


def item_font_size(card, width, height, max_size):
    """Mayor tamaño de letra (pt) con el que el texto cabe en la tarjeta."""
    width_pt, height_pt = width / 12700, height / 12700
    for size in range(max_size, 11, -1):
        chars_per_line = width_pt / (size * 0.47)
        lines = sum(math.ceil((len(a) + len(b) + 1) / chars_per_line) for a, b in card["items"])
        needed = lines * size * 1.2 * 1.08 + len(card["items"]) * size * 0.75
        if needed <= height_pt:
            return size
    return 12


def add_card(slide, x, y, w, h, card, n_cols, size):
    accent = card["color"]
    add_rect(slide, x, y, w, h, CARD_BG, MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.04
    add_rect(slide, x, y + Inches(0.12), Inches(0.07), h - Inches(0.24), accent)

    pad = Inches(0.3)
    inner_w = w - pad - Inches(0.25)
    add_textbox(slide, x + pad, y + Inches(0.25), inner_w, Inches(0.8),
                card["title"], 20 if n_cols == 2 else 18, accent, bold=True)

    items_top = Inches(1.15)
    tb = slide.shapes.add_textbox(x + pad, y + items_top, inner_w, h - items_top - Inches(0.2))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, (label, text) in enumerate(card["items"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(size * 0.75)
        p.line_spacing = 1.08
        r1 = p.add_run()
        style_run(r1, size, TEXT_WHITE, bold=True)
        r1.text = label + " "
        r2 = p.add_run()
        style_run(r2, size, TEXT_MUTED)
        r2.text = text


def build_slide(prs, data, index, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, BG_COLOR)
    sw, sh = prs.slide_width, prs.slide_height
    margin = Inches(0.6)

    first_color = data["cards"][0]["color"]
    add_textbox(slide, margin, Inches(0.45), Inches(9), Inches(0.35),
                data["category"], 13, first_color, bold=True)
    add_textbox(slide, margin, Inches(0.8), sw - 2 * margin, Inches(0.8),
                data["title"], 30 if len(data["title"]) <= 60 else 26, TEXT_WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(slide, margin, Inches(1.68), Inches(1.4), Inches(0.05), first_color)

    cards = data["cards"]
    n = len(cards)
    gap = Inches(0.3)
    top = Inches(2.0)
    bottom = sh - Inches(0.65)
    w = int((sw - 2 * margin - gap * (n - 1)) / n)
    # Mismo tamaño de letra en todas las tarjetas de la diapositiva
    inner_w, inner_h = w - Inches(0.55), bottom - top - Inches(1.35)
    size = min(item_font_size(c, inner_w, inner_h, 18 if n == 2 else 16) for c in cards)
    for i, card in enumerate(cards):
        add_card(slide, margin + i * (w + gap), top, w, bottom - top, card, n, size)

    add_textbox(slide, sw - margin - Inches(2), sh - Inches(0.45), Inches(2), Inches(0.3),
                f"{index} / {total}", 11, TEXT_MUTED).text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    add_textbox(slide, margin, sh - Inches(0.45), Inches(6), Inches(0.3),
                "Neurotransmisión sináptica", 11, TEXT_MUTED)

    slide.notes_slide.notes_text_frame.text = data["notes"]


def build_synaptic_presentation(filename="Neurotransmision_Sinaptica_Optimizada.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for i, data in enumerate(slides_data, start=1):
        build_slide(prs, data, i, len(slides_data))
    prs.save(filename)
    return filename


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "Neurotransmision_Sinaptica_Optimizada.pptx"
    print(build_synaptic_presentation(out))
